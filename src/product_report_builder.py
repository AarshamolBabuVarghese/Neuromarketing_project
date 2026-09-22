"""
product_report_builder.py
---------------------------------------------------------------
Rolls per-attribute-test results up into one consolidated,
plain-English recommendation per real product.

v4 CHANGE (this version): the CONTENT/LOGIC is unchanged from v3 --
same ADOPT/CONSIDER/STILL COLLECTING DATA classification, same
leaning/n_needed_estimate handling for inconclusive tests. What
changed is the VISUAL DESIGN of the .docx: it now uses the same
polished techniques as report_generator.py (colorful cover banner,
table of contents, shaded table headers, highlight boxes, embedded
matplotlib charts, page-number footer) instead of a plain
heading+bullets layout -- while still sourcing every number from
rec.py's correct, per-test-grouped results (never the confounded
global A/B comparison report_generator.py's OWN data source uses).

Install: pip install python-docx matplotlib

Drop into: Neuromarketing_project/src/product_report_builder.py
Run (after ab_recommendation_engine.py has produced results):
  python src/product_report_builder.py
---------------------------------------------------------------
"""

import os
import json
from collections import defaultdict
from typing import List, Dict, Any
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as patches

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import rec as engine

RESULTS_JSON = os.path.join(engine.OUT_DIR, "ab_recommendations.json")
REPORT_OUT_DIR = os.path.join(engine.OUT_DIR, "product_reports")
DOCX_OUT_PATH = os.path.join(REPORT_OUT_DIR, "AB_Test_Recommendations_Report.docx")
CSV_OUT_PATH = os.path.join(REPORT_OUT_DIR, "product_level_recommendations.csv")
CHART_DIR = os.path.join(REPORT_OUT_DIR, "_charts")

# ------------------------------------------------------------------
# Palette -- reused consistently across the cover banner, tables,
# headings and charts so the whole document reads as one design.
# ------------------------------------------------------------------
NAVY = RGBColor(0x1F, 0x2A, 0x44)
NAVY_HEX = "1F2A44"
GOLD = RGBColor(0xC0, 0x8A, 0x2E)
GREEN = RGBColor(0x1E, 0x7D, 0x32)
AMBER = RGBColor(0xB8, 0x86, 0x0B)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
CHART_PALETTE = ["#E84C6B", "#F5A623", "#7B5FE0", "#2FB6A6", "#1F2A44", "#C08A2E"]

ATTRIBUTE_ACTION_TEMPLATES = {
    "color": "Use {value} as the pack's primary color.",
    "color_and_branding": "Adopt the '{value}' color and branding treatment.",
    "packaging_format": "Package the product as: {value}.",
    "packaging_material": "Switch to {value} as the packaging material.",
    "packaging_transparency": "Use packaging with: {value}.",
    "packaging_finish": "Use a {value} on the pack.",
    "pack_format": "Use the '{value}' pack format.",
    "product_flavor": "Lead with the '{value}' variant.",
    "promotional_badge": "Feature: {value}.",
    "price_discount": "Price the product as: {value}.",
    "price_display": "Display pricing as: {value}.",
    "health_claims_badges": "Display these claims on-pack: {value}.",
    "packaging_claims": "Display these claims on-pack: {value}.",
    "product_imagery_vs_info": "Use this layout: {value}.",
    "product_imagery_layout": "Use this imagery layout: {value}.",
    "background_theme": "Use the '{value}' background theme.",
    "branding_design": "Use the '{value}' branding treatment.",
    "naming_and_size": "Use this naming/size: {value}.",
    "product_tagline": "Use the tagline: {value}.",
    "overall_packaging_redesign": "Adopt this overall packaging direction: {value}.",
}


def _sentence_for(attribute_changed: str, value: str) -> str:
    template = ATTRIBUTE_ACTION_TEMPLATES.get(attribute_changed, "Adopt: {value}.")
    return template.format(value=value)


# ------------------------------------------------------------------
# Content / logic -- UNCHANGED from v3
# ------------------------------------------------------------------
def load_results(json_path: str = RESULTS_JSON) -> List[Dict[str, Any]]:
    if not os.path.exists(json_path):
        raise FileNotFoundError(
            f"{json_path} not found. Run ab_recommendation_engine.py first."
        )
    with open(json_path, "r") as f:
        return json.load(f)


def group_by_product(results: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    grouped = defaultdict(list)
    for r in results:
        grouped[r.get("product_id", "UNKNOWN")].append(r)
    return grouped


def synthesize_product_recommendation(product_id: str, tests: List[Dict[str, Any]]) -> Dict[str, Any]:
    product_name = tests[0].get("product_name", "UNKNOWN")
    category = tests[0].get("category", "UNKNOWN")

    adopt, consider, inconclusive = [], [], []

    for t in tests:
        entry = {
            "attribute_changed": t.get("attribute_changed", "UNKNOWN"),
            "test_name": t["product_feature"],
            "confidence": t["composite_confidence"],
            "winner_variant": t["winner_variant"],
            "winner_value": t.get("winner_value"),
            "metric_results": t.get("metric_results", {}),
            "n_a": t.get("n_a", 0), "n_b": t.get("n_b", 0),
            "leaning_variant": t.get("leaning_variant"),
            "leaning_value": t.get("leaning_value"),
            "n_needed_estimate": t.get("n_needed_estimate"),
        }

        if t["winner_variant"] == "No clear winner":
            label = entry["attribute_changed"].replace("_", " ").title()
            if entry["leaning_value"]:
                sentence = (f"{label}: trending toward '{entry['leaning_value']}' "
                            f"(n_a={entry['n_a']}, n_b={entry['n_b']}) — not yet statistically "
                            f"significant, do not act on this yet.")
            elif entry["leaning_variant"]:
                sentence = (f"{label}: trending toward Variant {entry['leaning_variant']} "
                            f"(n_a={entry['n_a']}, n_b={entry['n_b']}) — not yet significant.")
            else:
                sentence = (f"{label}: no data or no directional signal yet "
                            f"(n_a={entry['n_a']}, n_b={entry['n_b']}).")
            if entry["n_needed_estimate"]:
                sentence += f" Collect ~{entry['n_needed_estimate']} more participants for a reliable read."
            entry["sentence"] = sentence
            inconclusive.append(entry)
            continue

        sentence = _sentence_for(entry["attribute_changed"], entry["winner_value"] or f"Variant {entry['winner_variant']}")
        entry["sentence"] = sentence

        if entry["confidence"] == "High":
            adopt.append(entry)
        else:
            consider.append(entry)

    total = len(tests)
    summary = (
        f"{product_name} — {total} attribute{'s' if total != 1 else ''} tested, "
        f"{len(adopt)} high-confidence change{'s' if len(adopt) != 1 else ''} recommended, "
        f"{len(consider)} worth considering with more data, "
        f"{len(inconclusive)} still collecting data."
    )

    return {
        "product_id": product_id, "product_name": product_name, "category": category,
        "summary": summary, "adopt": adopt, "consider": consider, "inconclusive": inconclusive,
    }


def build_all_product_recommendations(results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    grouped = group_by_product(results)
    recs = [synthesize_product_recommendation(pid, tests) for pid, tests in grouped.items()]
    recs.sort(key=lambda r: -len(r["adopt"]))
    return recs


# ==================================================================
# LOW-LEVEL DOCX STYLING HELPERS
# ==================================================================
def _shade_cell(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _header_row(table, headers):
    row = table.rows[0]
    for i, text in enumerate(headers):
        cell = row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(cell, NAVY_HEX)


def _add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _header_row(table, headers)
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return table


def _add_highlight_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _shade_cell(cell, "F3EFE6")
    cell.paragraphs[0].text = text
    for run in cell.paragraphs[0].runs:
        run.font.size = Pt(11)
    if not cell.paragraphs[0].runs:
        run = cell.paragraphs[0].add_run(text)
        run.font.size = Pt(11)
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:color"), NAVY_HEX)
        borders.append(el)
    tblPr.append(borders)


def _add_accent_bar(doc, text, color_hex):
    """A slim, colored full-width bar used as a section divider before a product heading."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _shade_cell(cell, color_hex)
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _add_toc(doc):
    doc.add_heading("Table of Contents", level=1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Right-click and choose 'Update Field' to populate the table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)
    doc.add_page_break()


def _add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


# ==================================================================
# CHART / COVER GRAPHIC GENERATION (matplotlib -> PNG -> embedded)
# ==================================================================
def _style_ax(ax):
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=9)


def _chart_cover_banner(out_path, width_in=8.5, height_in=2.6):
    """Abstract colorful geometric banner for the cover page, since no
    real design asset exists -- generated once per run so it doesn't
    require any external image file."""
    fig, ax = plt.subplots(figsize=(width_in, height_in))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    triangles = [
        ([0, 3, 6], [0, 3, 0], "#E84C6B"),
        ([2, 5, 8], [0, 3, 0], "#F5A623"),
        ([4, 7, 10], [0, 3, 0], "#7B5FE0"),
        ([6, 9, 10], [0, 2.4, 0], "#2FB6A6"),
    ]
    for xs, ys, color in triangles:
        ax.add_patch(patches.Polygon(list(zip(xs, ys)), closed=True, color=color, alpha=0.92, zorder=1))
    fig.patch.set_alpha(0)
    ax.patch.set_alpha(0)
    fig.tight_layout(pad=0)
    fig.savefig(out_path, dpi=200, transparent=True)
    plt.close(fig)


def _chart_confidence_pie(n_adopt, n_consider, n_collecting, out_path):
    labels, values, colors = [], [], []
    for label, val, color in [("High confidence", n_adopt, "#2FB6A6"),
                               ("Worth considering", n_consider, "#F5A623"),
                               ("Still collecting data", n_collecting, "#B8B8B8")]:
        if val > 0:
            labels.append(label)
            values.append(val)
            colors.append(color)
    fig, ax = plt.subplots(figsize=(5.2, 4))
    ax.pie(values, labels=labels, autopct="%1.0f%%", colors=colors, startangle=90,
           textprops={"fontsize": 9}, labeldistance=1.15)
    ax.set_title("Recommendation Readiness", fontsize=12, color=f"#{NAVY_HEX}")
    fig.subplots_adjust(left=0.15, right=0.85)
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def _chart_category_pie(product_recs, out_path):
    counts = defaultdict(int)
    for r in product_recs:
        counts[r["category"]] += 1
    fig, ax = plt.subplots(figsize=(4.6, 4))
    ax.pie(list(counts.values()), labels=[k.title() for k in counts.keys()],
           autopct="%1.0f%%", colors=CHART_PALETTE, startangle=90, textprops={"fontsize": 9})
    ax.set_title("Products Tested by Category", fontsize=12, color=f"#{NAVY_HEX}")
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def _chart_top_products_bar(product_recs, out_path, top_n=8):
    ranked = [r for r in product_recs if len(r["adopt"]) > 0][:top_n]
    if not ranked:
        return False
    names = [r["product_name"] for r in ranked][::-1]
    counts = [len(r["adopt"]) for r in ranked][::-1]
    fig, ax = plt.subplots(figsize=(7, max(2.5, 0.45 * len(names))))
    ax.barh(names, counts, color="#2FB6A6")
    ax.set_xlabel("High-confidence recommendations")
    ax.set_title("Top Products Ready to Act On", fontsize=12, color=f"#{NAVY_HEX}")
    _style_ax(ax)
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    return True


# ==================================================================
# MAIN DOCX BUILDER
# ==================================================================
def build_docx_report(product_recs: List[Dict[str, Any]], out_path: str = DOCX_OUT_PATH):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    os.makedirs(CHART_DIR, exist_ok=True)

    total_tests = sum(len(r["adopt"]) + len(r["consider"]) + len(r["inconclusive"]) for r in product_recs)
    total_adopt = sum(len(r["adopt"]) for r in product_recs)
    total_consider = sum(len(r["consider"]) for r in product_recs)
    total_inconclusive = sum(len(r["inconclusive"]) for r in product_recs)
    ready_tests = total_adopt + total_consider

    # ---- generate all chart images up front ----
    banner_path = os.path.join(CHART_DIR, "cover_banner.png")
    conf_pie_path = os.path.join(CHART_DIR, "confidence_pie.png")
    cat_pie_path = os.path.join(CHART_DIR, "category_pie.png")
    top_bar_path = os.path.join(CHART_DIR, "top_products.png")

    _chart_cover_banner(banner_path)
    _chart_confidence_pie(total_adopt, total_consider, total_inconclusive, conf_pie_path)
    _chart_category_pie(product_recs, cat_pie_path)
    has_top_chart = _chart_top_products_bar(product_recs, top_bar_path)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ---------- COVER PAGE ----------
    doc.add_picture(banner_path, width=Inches(6.5))
    pic_p = doc.paragraphs[-1]
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    title = doc.add_heading("A/B Neuromarketing Test", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Recommendations Report")
    run.font.size = Pt(18)
    run.font.color.rgb = GOLD
    run.italic = True

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"Generated {datetime.now().strftime('%B %d, %Y')}\n"
        f"{len(product_recs)} products analyzed  ·  {total_tests} attribute tests  ·  "
        f"{total_adopt} ready to adopt now"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = GREY
    doc.add_page_break()

    # ---------- TABLE OF CONTENTS ----------
    _add_toc(doc)

    # ---------- EXECUTIVE SUMMARY ----------
    doc.add_heading("Executive Summary", level=1)
    _add_highlight_box(
        doc,
        f"Of {total_tests} attribute tests across {len(product_recs)} products, "
        f"{total_adopt} are backed by strong enough evidence to adopt immediately, "
        f"{total_consider} show a promising signal worth validating further, and "
        f"{total_inconclusive} still need more participant data before any conclusion "
        f"can be drawn. A change is only marked ADOPT once at least two of the three "
        f"response metrics (engagement, emotion, attention time) show a statistically "
        f"significant difference (p < 0.05)."
    )
    doc.add_paragraph()

    doc.add_picture(conf_pie_path, width=Inches(3.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(cat_pie_path, width=Inches(3.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ---------- DATA COLLECTION STATUS ----------
    doc.add_heading("Data Collection Status", level=1)
    doc.add_paragraph(
        f"{ready_tests} of {total_tests} attribute tests have enough data to support a "
        f"recommendation so far. The remainder need more participant sessions — each test "
        f"needs roughly {engine.TARGET_N_PER_VARIANT} responses per variant for a reliable read."
    )
    doc.add_page_break()

    # ---------- TOP OPPORTUNITIES ----------
    if has_top_chart:
        doc.add_heading("Top Opportunities", level=1)
        doc.add_paragraph("Products with the most high-confidence, ready-to-implement changes:")
        doc.add_picture(top_bar_path, width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()

    # ---------- EXECUTIVE SUMMARY TABLE ----------
    doc.add_heading("All Products — Summary Table", level=1)
    rows = [[r["product_name"], r["category"].title(), len(r["adopt"]), len(r["consider"]), len(r["inconclusive"])]
            for r in product_recs]
    _add_table(doc, ["Product", "Category", "High-confidence", "Consider", "Collecting data"], rows,
               widths=[2.2, 1.3, 1.3, 1.1, 1.3])
    doc.add_page_break()

    # ---------- PER-PRODUCT RECOMMENDATIONS ----------
    doc.add_heading("Per-Product Recommendations", level=1)
    for r in product_recs:
        _add_accent_bar(doc, f"  {r['product_name']}   ·   {r['category'].title()}", NAVY_HEX)
        doc.add_paragraph(r["summary"]).italic = True

        _add_bullet_section(doc, "🟢 ADOPT — high confidence", GREEN, r["adopt"])
        _add_bullet_section(doc, "🟡 CONSIDER — medium confidence, worth more data", AMBER, r["consider"])
        _add_bullet_section(doc, "⚪ STILL COLLECTING DATA — directional only, not yet actionable", GREY, r["inconclusive"])
        doc.add_paragraph()

    _add_page_number_footer(doc)

    doc.save(out_path)
    print(f"Word report saved to {out_path}")


def _add_bullet_section(doc: Document, heading: str, color: RGBColor, entries: List[Dict[str, Any]]):
    if not entries:
        return
    p = doc.add_paragraph()
    run = p.add_run(heading)
    run.bold = True
    run.font.color.rgb = color
    for e in entries:
        bullet = doc.add_paragraph(style="List Bullet")
        bullet.add_run(f"{e['attribute_changed'].replace('_', ' ').title()}: ").bold = True
        bullet.add_run(e["sentence"])
        sig_bits = [f"{m.replace('_', ' ')} p={r['p_value']}"
                    for m, r in e.get("metric_results", {}).items() if r.get("significant")]
        if sig_bits:
            sub = doc.add_paragraph()
            sub.paragraph_format.left_indent = Inches(0.5)
            sub_run = sub.add_run("Support: " + "; ".join(sig_bits))
            sub_run.italic = True
            sub_run.font.size = Pt(9)
            sub_run.font.color.rgb = GREY


def export_product_csv(product_recs: List[Dict[str, Any]], out_path: str = CSV_OUT_PATH):
    import pandas as pd
    rows = []
    for r in product_recs:
        for bucket_name, bucket in (("adopt", r["adopt"]), ("consider", r["consider"]),
                                     ("collecting_data", r["inconclusive"])):
            for e in bucket:
                rows.append({
                    "product_id": r["product_id"], "product_name": r["product_name"],
                    "category": r["category"], "action_bucket": bucket_name,
                    "attribute_changed": e["attribute_changed"],
                    "recommended_value": e.get("winner_value") or e.get("leaning_value"),
                    "confidence": e.get("confidence", "N/A"),
                    "n_a": e.get("n_a"), "n_b": e.get("n_b"),
                    "sentence": e["sentence"],
                })
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    pd.DataFrame(rows).to_csv(out_path, index=False)
    print(f"Flat CSV saved to {out_path}")


if __name__ == "__main__":
    results = load_results()
    product_recs = build_all_product_recommendations(results)

    print(f"\nRolled {len(results)} attribute tests up into {len(product_recs)} products.\n")
    for r in product_recs:
        print(f"* {r['summary']}")

    build_docx_report(product_recs)
    export_product_csv(product_recs)